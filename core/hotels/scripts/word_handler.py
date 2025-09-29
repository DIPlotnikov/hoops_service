from docx import Document
from datetime import timedelta, datetime
from ..scripts import server_handler as SH
from shutil import copyfile
import os
from ..config import bucket as bucket
from ..config import core as core


class WordBuilder(object):
    dict_sample_names = {
        'CONTRACT': 'Договор.docx',

    }

    def __init__(self, *, path, type_report):
        self.path = path
        self.type_report = type_report

        self.sample_name = self.dict_sample_names.get(type_report)
        self.path = os.path.join(self.path, self.sample_name)

        self.__path_storage_prefix = f'/usr/local/share/minio/{bucket}'
        self.__path_sample_prefix = f'./{core}/scripts/sample'

        if os.path.isfile(os.path.join(self.__path_storage_prefix, self.path)):
            os.remove(os.path.join(self.__path_storage_prefix, self.path))

        os.makedirs(os.path.dirname(os.path.join(self.__path_storage_prefix, self.path)), exist_ok=True)
        copyfile(os.path.join(self.__path_sample_prefix, self.sample_name),
                 os.path.join(self.__path_storage_prefix, self.path))
        self.doc = Document(self.path_storage)

    @property
    def path_storage(self):
        return os.path.join(self.__path_storage_prefix, self.path)

    @property
    def path_sample(self):
        if self.sample_name:
            return os.path.join(self.__path_sample_prefix, self.sample_name)

    def pre_build(self):
        if os.path.isfile(self.path_storage):
            os.remove(self.path_storage)
        os.makedirs(os.path.dirname(self.path_storage), exist_ok=True)
        copyfile(self.path_sample, self.path_storage)

    def __replace(self, pattern, new_text):
        for p in self.doc.paragraphs:
            if pattern in p.text:
                p.text = p.text.replace(pattern, new_text)
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if pattern in p.text:
                            p.text = p.text.replace(pattern, new_text)

    def executer_offer(self, *, fio: str, passport: str, passport_object: str, date: datetime) -> str:
        if date is None:
            date = datetime.now()
        self.__replace('{FIO}', fio)
        self.__replace('{DATE}',f'{str(date.day).zfill(2)}.{str(date.month).zfill(2)}.{date.year}')
        self.__replace('{PASSPORT}', passport)
        self.__replace('{PASSPORT_OBJECT}', passport_object)
        self.doc.save(self.path_storage)
        return self.path
